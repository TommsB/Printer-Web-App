"""One-off import of the existing consumables Excel into the SQLite DB.

Usage: python scripts/import_excel.py <file.xlsx> [sheet_name]

Expected layout (as in the current workbook):
  Table 1 (printers):  A=IP, B=Lokācija, C=Modelis, D=Ražotājs, E=(ignored),
                       F=Krāsains/Melnbalts, G..J=toner codes
  Table 2 (stock):     header row containing "IP Adrese", "Lokācija", "Toneris",
                       "Rezerve (skaits)", "Optimālā rezerve"; blank IP/location cells
                       inherit the value from the row above (merged cells).
Company is derived from the location prefix (TXP / TP / TX) — edit COMPANY_BY_PREFIX.
"""

import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.db import get_db, init_db  # noqa: E402

COMPANY_BY_PREFIX = {"TXP": "Tenax Panel", "TP": "Tenapors", "TX": "Tenax"}
IP_RE = re.compile(r"^\d{1,3}(\.\d{1,3}){3}$")
COLORS = {"K": "K", "C": "C", "M": "M", "Y": "Y"}


def company_for(location: str) -> str:
    return COMPANY_BY_PREFIX.get(location.split(" ")[0].upper(), "")


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return str(v).strip() if v is not None else ""


def main(path: str, sheet: str | None) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    init_db()
    printers = toners = links = stock_rows = 0

    with get_db() as conn:
        def toner_id(code: str, color: str = "") -> int:
            nonlocal toners
            row = conn.execute("SELECT id FROM toner_models WHERE code = ?", (code,)).fetchone()
            if row:
                return row["id"]
            toners += 1
            return conn.execute("INSERT INTO toner_models (code, color) VALUES (?,?)",
                                (code, color)).lastrowid

        # Toner columns are headed "Toner K/C/Y/M" on row 1.
        def header_color(c: int) -> str:
            return COLORS.get(cell(ws, 1, c)[-1:].upper(), "")

        stock_header = None
        for r in range(1, ws.max_row + 1):
            a = cell(ws, r, 1)
            if a == "IP Adrese" and cell(ws, r, 3) == "Toneris":
                stock_header = r
                break
            if not IP_RE.match(a):
                continue
            location = cell(ws, r, 2)
            conn.execute(
                "INSERT INTO printers (company, location, model, brand, ip, color_type) VALUES (?,?,?,?,?,?)"
                " ON CONFLICT(ip) DO UPDATE SET company=excluded.company, location=excluded.location,"
                " model=excluded.model, brand=excluded.brand, color_type=excluded.color_type",
                (company_for(location), location, cell(ws, r, 3), cell(ws, r, 4), a, cell(ws, r, 6)))
            pid = conn.execute("SELECT id FROM printers WHERE ip = ?", (a,)).fetchone()["id"]
            printers += 1
            for c in range(7, 11):
                code = cell(ws, r, c)
                if code:
                    conn.execute("INSERT OR IGNORE INTO printer_toners (printer_id, toner_id) VALUES (?,?)",
                                 (pid, toner_id(code, header_color(c))))
                    links += 1

        if stock_header:
            ip = ""
            for r in range(stock_header + 1, ws.max_row + 1):
                ip = cell(ws, r, 1) or ip
                code = cell(ws, r, 3)
                if not code or not IP_RE.match(ip):
                    continue
                p = conn.execute("SELECT id FROM printers WHERE ip = ?", (ip,)).fetchone()
                if not p:
                    print(f"  skip stock row {r}: unknown printer {ip}")
                    continue
                tid = toner_id(code)

                def num(c):
                    v = ws.cell(r, c).value
                    return int(v) if isinstance(v, (int, float)) else 0

                conn.execute(
                    "INSERT INTO printer_toners (printer_id, toner_id, qty, optimal_qty) VALUES (?,?,?,?)"
                    " ON CONFLICT(printer_id, toner_id) DO UPDATE SET qty=excluded.qty, optimal_qty=excluded.optimal_qty",
                    (p["id"], tid, num(4), num(5)))
                stock_rows += 1

    print(f"Printers: {printers}, new toner models: {toners}, links: {links}, stock rows: {stock_rows}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
